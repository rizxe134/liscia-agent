"""v1.3.0 planner regression: 40+ phrasings, model mocked (rules-only, garbage, and plausible-wrong model output)."""
import json, os, sys, unittest
from unittest import mock
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import planner, safety

MIL = "Create a Word document and an Excel sheet listing the top 10 most powerful militaries in the world based on recent rankings or reports."
CASES = [  # (task, intent, outputs)
    (MIL, "web_research_to_files", ["docx", "xlsx"]),
    ("can you open up youtube and on youtube can you search for the song Pursuit by isq and claick on to play the song", "media_play", []),
    ("play Pursuit by isq on youtube", "media_play", []),
    ("open youtube and play lofi hip hop", "media_play", []),
    ("youtube despacito", "media_play", []),
    ("put on some taylor swift on youtube", "media_play", []),
    ("open hulu", "open_site", []),
    ("go to hulu.com", "open_site", []),
    ("open netflix", "open_site", []),
    ("take me to amazon", "open_site", []),
    ("open reddit", "open_site", []),
    ("Make a Word document of the top 20 richest people in the world from Forbes", "web_research_to_files", ["docx"]),
    ("create a word doc listing the 20 richest people according to forbes", "web_research_to_files", ["docx"]),
    ("Open Excel and make a sheet of the top 20 richest people in the world", "web_research_to_files", ["xlsx"]),
    ("open excel and list the 10 largest countries by area", "web_research_to_files", ["xlsx"]),
    ("make a spreadsheet of the 5 biggest companies by market cap", "web_research_to_files", ["xlsx"]),
    ("put the 10 tallest buildings in an excel file", "web_research_to_files", ["xlsx"]),
    ("write a word document about the 5 most populous cities", "web_research_to_files", ["docx"]),
    ("make a doc and a spreadsheet of the top 10 best selling cars of 2025", "web_research_to_files", ["docx", "xlsx"]),
    ("save the top 5 fastest cars to a text file", "web_research_to_files", ["txt"]),
    ("excel sheet with top 10 most powerful militaries", "web_research_to_files", ["xlsx"]),
    ("I need a word file and an excel table of the 10 richest countries by GDP", "web_research_to_files", ["docx", "xlsx"]),
    ("google who is the richest person in the world right now", "web_answer", []),
    ("search google for the richest person in the world", "web_answer", []),
    ("who is the richest person in the world", "web_answer", []),
    ("what is the capital of australia", "web_answer", []),
    ("who won the last world cup", "web_answer", []),
    ("how tall is mount everest", "web_answer", []),
    ("what's the weather in new york today", "web_answer", []),
    ("when is the next spacex launch", "web_answer", []),
    ("open calculator", "desktop", []),
    ("open notepad and type hello world", "desktop", []),
    ("open calculator and add 2+2", "desktop", []),
    ("open file explorer", "desktop", []),
    ("open spotify app", "desktop", []),
    ("take a screenshot of my desktop", "desktop", []),
    ("open the settings app and turn on dark mode", "desktop", []),
    ("open amazon and add a usb cable to the cart", "web_interact", []),
    ("go on the page and look for daft punk songs", "web_interact", []),
    ("log into my gmail and check unread emails", "web_interact", []),
    ("open youtube and click the first video", "media_play", []),
    ("search wikipedia for alan turing", None, []),
]
OK_ALT = {"search wikipedia for alan turing": {"web_answer", "web_interact", "open_site"},
          "log into my gmail and check unread emails": {"web_interact", "open_site"},
          "open amazon and add a usb cable to the cart": {"web_interact", "open_site"},
          "go on the page and look for daft punk songs": {"web_interact", "web_answer"},
          "take a screenshot of my desktop": {"desktop"},
          "what's the weather in new york today": {"web_answer"}}

def model(reply):
    c = mock.MagicMock(); c.text_chat.return_value = reply; return c

class PlannerPhrasings(unittest.TestCase):
    def check(self, p, task, intent, outs):
        ok = OK_ALT.get(task, {intent})
        self.assertIn(p["intent"], ok, f"{task!r} -> {p}")
        self.assertEqual(sorted(p.get("outputs") or []), sorted(outs), f"{task!r} -> {p}")

    def test_at_least_40(self): self.assertGreaterEqual(len(CASES), 40)

    def test_rules_only(self):
        for task, intent, outs in CASES:
            with self.subTest(task=task): self.check(planner.plan(None, task, use_model=False), task, intent, outs)

    def test_correct_model(self):
        for task, intent, outs in CASES:
            want = intent or "web_answer"
            c = model(json.dumps({"intent": want, "query": task, "outputs": outs, "count": 0}))
            with self.subTest(task=task): self.check(planner.plan(c, task), task, intent, outs)

    def test_garbage_model_falls_back(self):
        for reply in ["", "not json at all", "{", "[]", '{"intent":"dance","outputs":"x"}', None]:
            for task, intent, outs in CASES:
                c = model(reply)
                with self.subTest(task=task, reply=reply): self.check(planner.plan(c, task), task, intent, outs)

    def test_model_exception_falls_back(self):
        c = mock.MagicMock(); c.text_chat.side_effect = RuntimeError("ollama down")
        p = planner.plan(c, MIL); self.assertEqual(p["intent"], "web_research_to_files"); self.assertEqual(p["outputs"], ["docx", "xlsx"])

    def test_plausible_wrong_model_corrected(self):
        # the exact v1.2.8 failure: model says desktop (drive Word/Excel by vision) -> code forces files
        for wrong in [{"intent": "desktop", "query": "", "outputs": []},
                      {"intent": "web_answer", "query": "top militaries", "outputs": []},
                      {"intent": "web_research_to_files", "query": "militaries", "outputs": ["docx"], "count": 3}]:
            p = planner.plan(model(json.dumps(wrong)), MIL)
            self.assertEqual(p["intent"], "web_research_to_files"); self.assertEqual(sorted(p["outputs"]), ["docx", "xlsx"]); self.assertEqual(p["count"], 10)
        p = planner.plan(model('{"intent":"desktop","query":"x","outputs":[]}'), "who is the richest person in the world")
        self.assertNotEqual(p["intent"], "desktop")

    def test_militaries_query(self):
        p = planner.plan(None, MIL, use_model=False)
        self.assertEqual(p["count"], 10)
        self.assertIn(str(planner.today().year), p["query"])
        self.assertNotIn("based on", p["query"].lower()); self.assertNotIn("reports", p["query"].lower())
        self.assertIn("militar", p["query"].lower())

    def test_counts(self):
        for task, n in [("top 5 fastest cars in a word doc", 5), ("make an excel sheet of the twenty richest people", 20), (MIL, 10)]:
            self.assertEqual(planner.plan(None, task, use_model=False)["count"], n, task)

    def test_unambiguous_skips_model(self):
        for t in ["open hulu", "play Pursuit by isq on youtube", "open youtube and click the first video"]:
            c = model("{}"); planner.plan(c, t); c.text_chat.assert_not_called()

class VisionBlocklist(unittest.TestCase):
    def test_blocks(self):
        t = "Create a Word document and an Excel sheet listing the top 10 militaries"
        for act in [{"action": "hotkey", "keys": ["win", "r"]}, {"action": "hotkey", "keys": "winleft+r"},
                    {"action": "type", "text": "services.msc"}, {"action": "open_app", "app": "regedit"},
                    {"action": "open_app", "app": "cmd"}, {"action": "open_app", "app": "PowerShell"},
                    {"action": "open_app", "app": "Task Manager"}, {"action": "hotkey", "keys": ["ctrl", "shift", "esc"]},
                    {"action": "open_app", "app": "Settings"}, {"action": "open_app", "app": "Microsoft Word"},
                    {"action": "open_app", "app": "Excel"}]:
            self.assertTrue(safety.vision_block(act, t), act)
    def test_allows_when_asked_or_harmless(self):
        self.assertIsNone(safety.vision_block({"action": "open_app", "app": "Settings"}, "open the settings app and turn on dark mode"))
        self.assertIsNone(safety.vision_block({"action": "open_app", "app": "Task Manager"}, "open task manager"))
        self.assertIsNone(safety.vision_block({"action": "hotkey", "keys": ["win", "r"]}, "press win+r and type notepad"))
        self.assertIsNone(safety.vision_block({"action": "open_app", "app": "Calculator"}, "open calculator"))
        self.assertIsNone(safety.vision_block({"action": "type", "text": "hello world"}, "open notepad and type hello world"))
        self.assertIsNone(safety.vision_block({"action": "hotkey", "keys": ["ctrl", "c"]}, "copy"))

if __name__ == "__main__": unittest.main()


CARS = "can you look up what the top 10 fastest cars are then make make a list on exacel hseet as well as a report on word with an addtion of a small power point of it"

class V131ChatVsTask(unittest.TestCase):
    def setUp(self):
        import core; self.core = core
    def test_router_cases(self):
        client = mock.MagicMock(); client.text_chat.return_value = '{"type":"chat"}'
        tasks = [CARS, "make a powerpoint about the 5 largest oceans", "put the top 5 phones in a pwerpoint",
                 "search for cheap flights to paris", "find the best laptops of 2026 and save them to an exel file",
                 "look up the 10 tallest buildings and make a wrod doc", "create a slide deck of the top 5 cities by population",
                 "open youtube and play lofi"]
        for m in tasks:
            kind, out = self.core.route(client, m)
            self.assertEqual(kind, "task", m); self.assertEqual(out, m)
        for m in ["hi how are you", "what is linear algebra", "thanks!", "tell me a joke", "how do I make an excel chart"]:
            kind, _ = self.core.route(client, m)
            self.assertEqual(kind, "chat", m)
    def test_chat_prompt_says_it_can(self):
        self.assertIn("CAN", self.core.CHAT_SYSTEM); self.assertIn("PowerPoint", self.core.CHAT_SYSTEM)
    def test_cars_plan(self):
        p = planner.plan(model('{"intent":"desktop","query":"x","outputs":[]}'), CARS)
        self.assertEqual(p["intent"], "web_research_to_files"); self.assertEqual(sorted(p["outputs"]), ["docx", "pptx", "xlsx"])
        self.assertEqual(p["count"], 10); self.assertEqual(p["query"], "top 10 fastest cars")
    def test_typo_outputs(self):
        for t, k in [("exacel", "xlsx"), ("exel", "xlsx"), ("wrod", "docx"), ("power point", "pptx"), ("pwerpoint", "pptx"),
                     ("ppt", "pptx"), ("slides", "pptx"), ("deck", "pptx"), ("presentation", "pptx"), ("hseet", "xlsx")]:
            self.assertIn(k, planner.file_outputs(f"top 5 cars in a {t}"), t)
    def test_query_cleaning(self):
        self.assertEqual(planner.strip_instructions("Open Microsoft Excel and create a new sheet. List the top 10 most powerful militaries"), "top 10 most powerful militaries")

class V131Files(unittest.TestCase):
    def test_pptx_and_report(self):
        import tempfile, docmaker
        from pptx import Presentation
        from docx import Document
        d = tempfile.mkdtemp(); rows = [[str(i), f"Car {i}", f"{300 + i} mph"] for i in range(1, 11)]
        p = docmaker.make("pptx", "Top 10 Fastest Cars", "Summary text", ["Rank", "Car", "Top speed"], rows, source="https://x.test", folder=d)
        slides = Presentation(p).slides
        self.assertEqual(len(slides), 1 + 1 + 2 + 1)  # title, summary, 10 rows / 6 per slide, source
        rep = {"intro": "Intro.", "items": [{"name": f"Car {i}", "text": "Fast."} for i in range(1, 11)], "conclusion": "End."}
        w = docmaker.make("docx", "Top 10 Fastest Cars", "S", ["Rank", "Car", "Top speed"], rows, folder=d, report=rep)
        text = "\n".join(x.text for x in Document(w).paragraphs)
        for s in ["Introduction", "Intro.", "Details", "Car 10", "Conclusion"]: self.assertIn(s, text)
