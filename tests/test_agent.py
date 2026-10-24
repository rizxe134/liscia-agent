import os, sys, json, unittest
from unittest import mock
from PIL import Image
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from actions import parse_action, ActionError, Executor, is_risky, search_url
import screen
from core import Agent
import actions as _actions, core as _core, intents as _intents
# Tests must never touch the real browser/OS: run them as "not macOS" unless a test patches it on purpose.
_actions.IS_MAC = False; _core.IS_MAC = False
_REAL_FRONT, _REAL_RECENT = _intents.browser_dom.frontmost_app, _intents.browser_dom.recent_browser
_intents.browser_dom.frontmost_app = lambda *a, **k: None
_intents.browser_dom.recent_browser = lambda *a, **k: None
_intents.GUESS_CHECK = False

class Parse(unittest.TestCase):
    def test_click_in_noise(self):
        a = parse_action('Sure! ```json\n{"action":"click","x":"10.6","y":20}\n```')
        self.assertEqual((a["action"], a["x"], a["y"]), ("click", 11, 20))
    def test_hotkey_string(self):
        self.assertEqual(parse_action('{"action":"hotkey","keys":"Ctrl+L"}')["keys"], ["ctrl", "l"])
    def test_errors(self):
        for bad in ["no json", '{"action":"fly"}', '{"action":"click","x":1}', '{"action":"type"}',
                    '{"action":"open_url"}', '{"action":"web_search"}', '{"action":"hotkey","keys":[]}']:
            with self.assertRaises(ActionError, msg=bad): parse_action(bad)
    def test_wait_clamped(self):
        self.assertEqual(parse_action('{"action":"wait","seconds":999}')["seconds"], 30.0)

class Execute(unittest.TestCase):
    def setUp(self):
        p = mock.patch("actions.IS_MAC", False); p.start(); self.addCleanup(p.stop)
        self.g, self.b = mock.MagicMock(), mock.MagicMock()
        self.ex = Executor(self.g, scale=2.0, move_duration=0, browser=self.b, sleep=mock.MagicMock())
    def test_click_scaled(self):
        self.ex.execute(parse_action('{"action":"click","x":100,"y":50}'))
        self.g.click.assert_called_once_with(200, 100)
    def test_type_enter(self):
        self.ex.execute(parse_action('{"action":"type","text":"hi","press_enter":true}'))
        self.g.write.assert_called_once(); self.g.press.assert_called_once_with("enter")
    def test_hotkey_scroll(self):
        self.ex.execute(parse_action('{"action":"hotkey","keys":["ctrl","t"]}')); self.g.hotkey.assert_called_with("ctrl", "t")
        self.ex.execute(parse_action('{"action":"scroll","amount":-3}')); self.g.scroll.assert_called_with(-300)
    def test_browser(self):
        self.ex.execute(parse_action('{"action":"open_url","url":"example.com"}'))
        self.b.open.assert_called_with("https://example.com")
        self.ex.execute(parse_action('{"action":"web_search","query":"ollama vision"}'))
        self.b.open.assert_called_with(search_url("ollama vision"))
    def test_done(self):
        self.assertTrue(self.ex.execute(parse_action('{"action":"done"}')))

class Risk(unittest.TestCase):
    W = ["buy", "send", "submit"]
    def test_cases(self):
        self.assertTrue(is_risky(parse_action('{"action":"click","x":1,"y":1,"reason":"click Buy now"}'), self.W))
        self.assertTrue(is_risky(parse_action('{"action":"type","text":"x","sensitive":true}'), self.W))
        self.assertFalse(is_risky(parse_action('{"action":"click","x":1,"y":1,"reason":"open search box"}'), self.W))

class Screen(unittest.TestCase):
    def test_prepare(self):
        img, s = screen.prepare(Image.new("RGB", (2560, 1440)), 1280, 100)
        self.assertEqual(img.size, (1280, 720)); self.assertEqual(s, 2.0)

class Loop(unittest.TestCase):
    def test_run(self, tmp="/tmp/pcagent-test"):
        cfg = {"ollama": {"host": "x", "model": "m"},
               "agent": {"max_steps": 5, "runs_dir": tmp, "plan": False, "step_delay_s": 0, "move_duration_s": 0},
               "safety": {"confirm_risky": True, "risky_words": ["send"], "kill_keys": []}}
        client = mock.MagicMock(); client.chat.side_effect = [
            "garbage", '{"action":"click","x":5,"y":5,"reason":"press Send"}',
            '{"action":"web_search","query":"q"}', '{"action":"done","summary":"ok"}']
        g = mock.MagicMock(); g.size.return_value = (800, 600); confirm = mock.MagicMock(return_value=False)
        with mock.patch("actions.webbrowser") as wb, mock.patch("actions.IS_MAC", False):
            agent = Agent(cfg, confirm=confirm, log=lambda m: None, gui=g,
                          capture=lambda: Image.new("RGB", (800, 600)), client=client)
            agent.kill.start = lambda: None
            self.assertEqual(agent.run("t"), "ok")
            wb.open.assert_called_once()
        confirm.assert_called_once(); g.click.assert_not_called()

if __name__ == "__main__":
    unittest.main()

class Route(unittest.TestCase):
    def test_route(self):
        from core import route
        c = mock.MagicMock()
        self.assertEqual(route(c, "/do open gmail"), ("task", "open gmail")); c.text_chat.assert_not_called()
        c.text_chat.return_value = '{"type":"chat","reply":"Hi!"}'
        self.assertEqual(route(c, "hello"), ("chat", "Hi!"))
        c.text_chat.return_value = '{"type":"task","task":"search weather"}'
        self.assertEqual(route(c, "check weather"), ("task", "check weather"))  # 1.3.1: the user's own words are kept
        self.assertEqual(route(c, "do it again"), ("task", "search weather"))
        c.text_chat.return_value = "plain text"
        self.assertEqual(route(c, "hey"), ("chat", "plain text"))
        c.text_chat.return_value = "sure"
        self.assertEqual(route(c, "/chat hi"), ("chat", "sure"))

class MacSupport(unittest.TestCase):
    def test_keys(self):
        from actions import norm_key
        self.assertEqual(norm_key("cmd", mac=True), "command"); self.assertEqual(norm_key("cmd", mac=False), "ctrl")
        self.assertEqual(norm_key("ctrl", mac=True), "ctrl"); self.assertEqual(norm_key("option", mac=False), "alt")
    def test_retina_ratio(self):
        g = mock.MagicMock(); g.size.return_value = (1512, 982)
        self.assertAlmostEqual(screen.point_ratio(g, 3024), 0.5)
        g.size.return_value = (1920, 1080); self.assertEqual(screen.point_ratio(g, 1920), 1.0)
    def test_retina_click(self):
        # 3024px retina screenshot -> 1280 image (scale 2.3625) -> points (x0.5)
        img, s = screen.prepare(Image.new("RGB", (3024, 1964)), 1280, 0)
        g = mock.MagicMock(); g.size.return_value = (1512, 982)
        ex = Executor(g, scale=s * screen.point_ratio(g, 3024), move_duration=0)
        ex.execute(parse_action('{"action":"click","x":640,"y":400}'))
        g.click.assert_called_once_with(756, 472)

class KillSwitchTest(unittest.TestCase):
    def test_no_pynput_on_mac(self):
        from safety import KillSwitch
        k = KillSwitch()
        with mock.patch("sys.platform", "darwin"), mock.patch.object(KillSwitch, "_start_mac") as m:
            k.start(); m.assert_called_once()
        k.event.set(); self.assertTrue(k.killed); k.stop()

class LoopDetection(unittest.TestCase):
    def test_guard_levels(self):
        from core import LoopGuard
        g = LoopGuard(2, 3, 5); a = parse_action('{"action":"type","text":"youtube.com"}')
        self.assertEqual([g.check(a) for _ in range(5)], [0, 1, 2, 2, 3])
        self.assertEqual(g.check(parse_action('{"action":"open_app","app":"Safari"}')), 0)
    def test_nearby_clicks_same(self):
        from core import action_signature
        a = parse_action('{"action":"click","x":101,"y":201}'); b = parse_action('{"action":"click","x":105,"y":205}')
        self.assertEqual(action_signature(a), action_signature(b))

    def _cfg(self, tmp):
        return {"ollama": {"host": "x", "model": "m"},
                "agent": {"max_steps": 10, "runs_dir": tmp, "step_delay_s": 0, "plan": True},
                "safety": {"confirm_risky": False, "kill_keys": []}}

    def test_repeating_model_is_stopped_and_warned(self):
        client = mock.MagicMock(); client.text_chat.return_value = "1. open Safari"
        client.chat.return_value = '{"action":"type","text":"youtube.com"}'
        ex = mock.MagicMock(); g = mock.MagicMock(); g.size.return_value = (800, 600)
        agent = Agent(self._cfg("/tmp/pcagent-test"), log=lambda m: None, gui=g, client=client,
                      capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: ex)
        agent.kill.start = lambda: None
        res = agent.run("t")
        self.assertIn("stuck repeating", res)
        self.assertEqual(ex.execute.call_count, 2)          # repeats 3 and 4 were blocked
        prompts = [c.args[1] for c in client.chat.call_args_list]
        self.assertIn("PLAN:", prompts[0])
        self.assertIn("NO VISIBLE CHANGE", prompts[1])       # static screen detected
        self.assertIn("WARNING", prompts[2]); self.assertIn("NOT working", prompts[3])

    def test_screen_change_detected(self):
        frames = iter([Image.new("RGB", (800, 600)), Image.new("RGB", (800, 600), "white"), Image.new("RGB", (800, 600), "white")])
        client = mock.MagicMock(); client.chat.side_effect = ['{"action":"open_app","app":"Safari"}', '{"action":"done"}']
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = self._cfg("/tmp/pcagent-test"); cfg["agent"]["plan"] = False
        agent = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: next(frames), executor=lambda s: mock.MagicMock())
        agent.kill.start = lambda: None; agent.run("t")
        self.assertIn("screen changed", client.chat.call_args_list[1].args[1])

class HighLevelActions(unittest.TestCase):
    def test_open_app_and_browser(self):
        from actions import app_command
        self.assertEqual(app_command("Safari", "https://youtube.com", mac=True), ["open", "-a", "Safari", "https://youtube.com"])
        self.assertEqual(app_command("chrome", mac=False, windows=True), ["cmd", "/c", "start", "", "chrome"])
        run = mock.MagicMock(); b = mock.MagicMock()
        mock.patch("actions.IS_MAC", True).start(); self.addCleanup(mock.patch.stopall)
        ex = Executor(mock.MagicMock(), browser=b, runner=run, sleep=lambda s: None)
        ex.execute(parse_action('{"action":"open_app","app":"Safari"}')); self.assertEqual(run.call_args.args[0][-1], "Safari")
        ex.execute(parse_action('{"action":"open_url","url":"youtube.com","browser":"Safari"}'))
        self.assertEqual(run.call_args.args[0][-2:], ["Safari", "https://youtube.com"]); b.open.assert_not_called()
        with self.assertRaises(ActionError): parse_action('{"action":"open_app"}')
        self.assertTrue(parse_action('{"action":"type","text":"a","press_enter":"true"}')["press_enter"])

class BrowserDomTest(unittest.TestCase):
    def _runner(self, js_out):
        calls = []
        def run(cmd, **kw):
            calls.append(cmd); r = mock.MagicMock(returncode=0, stderr="")
            if cmd[:2] == ["lsappinfo", "front"]: r.stdout = "ASN:0x0-0x1234:"
            elif cmd[0] == "lsappinfo": r.stdout = '"LSDisplayName"="Safari"'
            else: r.stdout = js_out
            return r
        return run, calls
    def test_snapshot_and_click(self):
        import browser_dom
        snap = {"url": "https://youtube.com", "title": "YouTube", "elements": [{"id": 0, "tag": "a", "type": "", "text": "Lofi video", "x": 1, "y": 2}]}
        run, calls = self._runner(json.dumps(snap))
        with mock.patch("browser_dom.sys.platform", "darwin"), mock.patch.object(browser_dom, "frontmost_app", _REAL_FRONT):
            d = browser_dom.BrowserDOM(run=run, log=lambda m: None); s = d.snapshot()
        self.assertEqual(s["app"], "Safari"); self.assertIn('[0] a "Lofi video"', browser_dom.format_elements(s))
        self.assertIn("do JavaScript", calls[-1][4])
        run2, calls2 = self._runner("ok"); d.run = run2; d.click(0); self.assertIn("__pca", calls2[-1][-1])
        ex = Executor(mock.MagicMock(), dom=d); ex.execute(parse_action('{"action":"click_element","id":"0"}'))
    def test_not_browser(self):
        import browser_dom
        run, _ = self._runner("")
        def run_finder(cmd, **kw):
            r = run(cmd, **kw)
            if cmd[0] == "lsappinfo" and cmd[1] == "info": r.stdout = '"LSDisplayName"="Finder"'
            return r
        with mock.patch("browser_dom.sys.platform", "darwin"), mock.patch.object(browser_dom, "frontmost_app", _REAL_FRONT):
            self.assertIsNone(browser_dom.BrowserDOM(run=run_finder, log=lambda m: None).snapshot())

class FastPathTest(unittest.TestCase):
    def test_dom_text_path_used(self):
        dom = mock.MagicMock(); dom.snapshot.return_value = {"app": "Safari", "url": "u", "title": "t",
            "elements": [{"id": 3, "tag": "a", "type": "", "text": "First video", "x": 1, "y": 1}]}
        client = mock.MagicMock(); client.num_predict = 200
        client.text_chat.side_effect = ['{"action":"click_element","id":3}', '{"action":"done","summary":"ok"}']
        ex = mock.MagicMock(); g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 5, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False, "planner": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        agent = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)),
                      executor=lambda s: ex, dom=dom)
        agent.kill.start = lambda: None
        self.assertEqual(agent.run("pick the cat clip from the list"), "ok")
        client.chat.assert_not_called()   # no vision call needed
        self.assertIn('[3] a "First video"', client.text_chat.call_args_list[0].args[0][1]["content"])
        self.assertEqual(ex.execute.call_args.args[0]["id"], 3)

class StreamingTest(unittest.TestCase):
    def test_stream_tokens(self):
        from llm import OllamaClient
        lines = [json.dumps({"message": {"content": p}, "done": False}).encode() for p in ("Hel", "lo")] + [json.dumps({"done": True}).encode()]
        resp = mock.MagicMock(status_code=200); resp.iter_lines.return_value = lines
        resp.__enter__ = lambda s: resp; resp.__exit__ = lambda *a: False
        got = []
        with mock.patch("llm.requests.post", return_value=resp) as post:
            out = OllamaClient("http://h", "m", text_model="t").text_chat([{"role": "user", "content": "hi"}], on_token=got.append)
        self.assertEqual((out, got), ("Hello", ["Hel", "lo"]))
        body = post.call_args.kwargs["json"]
        self.assertEqual((body["model"], body["keep_alive"], body["stream"], body["think"]), ("t", "30m", True, False))
    def test_route_streams_chat(self):
        from core import route
        c = mock.MagicMock(); c.text_chat.side_effect = ['{"type":"chat"}', "Hi there"]
        self.assertEqual(route(c, "hello", on_token=lambda p: None), ("chat", "Hi there"))
        self.assertIsNotNone(c.text_chat.call_args.kwargs["on_token"])

class NestedParams(unittest.TestCase):
    def test_flatten(self):
        a = parse_action('{"action":"open_app","params":{"app":"Safari"}}')
        self.assertEqual(a["app"], "Safari")

class TrailingJunk(unittest.TestCase):
    def test_two_objects(self):
        self.assertEqual(parse_action('{"action":"wait","seconds":1}\n{"action":"done"}')["action"], "wait")

class NormCoords(unittest.TestCase):
    def test_qwen3_norm(self):
        from core import uses_norm_coords, denormalize, NORM_SYSTEM
        self.assertTrue(uses_norm_coords({"ollama": {"model": "qwen3-vl:4b-instruct"}, "agent": {}}))
        self.assertFalse(uses_norm_coords({"ollama": {"model": "qwen2.5vl:7b"}, "agent": {}}))
        self.assertEqual(denormalize({"action": "click", "x": 138, "y": 960}, 1024, 665)["x"], 141)
        self.assertIn("NORMALIZED 0-1000", NORM_SYSTEM)
    def test_loop_denormalizes(self):
        client = mock.MagicMock(); client.chat.side_effect = ['{"action":"click","x":500,"y":500}', '{"action":"done"}']
        ex = mock.MagicMock(); g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "qwen3-vl:4b-instruct"}, "agent": {"max_steps": 3, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        agent = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: ex, dom=False)
        agent.kill.start = lambda: None; agent.run("t")
        a = ex.execute.call_args.args[0]; self.assertEqual((a["x"], a["y"]), (400, 300))
        self.assertIn("NORMALIZED", client.chat.call_args_list[0].args[0])

class IntentsTest(unittest.TestCase):
    def setUp(self):
        import intents; intents.STATE.update(browser=None, app=None)
    def test_plan(self):
        import intents
        acts, rest = intents.plan("open safari, search for youtube.com, and click on the first video")
        self.assertEqual([a["action"] for a in acts], ["open_app", "open_url", "first_video"])
        self.assertEqual((acts[0]["app"], acts[1]["url"], rest), ("Safari", "https://www.youtube.com", ""))
        self.assertEqual(intents.plan("open youtube")[0][0]["url"], "https://www.youtube.com")
        self.assertEqual(intents.plan("go to youtube.com in chrome")[0][0]["browser"], "Google Chrome")
        self.assertEqual(intents.plan("write an email to bob"), ([], "write an email to bob"))
        acts, rest = intents.plan("open safari and write a poem in notes")
        self.assertEqual((len(acts), rest), (1, "write a poem in notes"))
    def test_router_imperative(self):
        from core import route
        c = mock.MagicMock()
        self.assertEqual(route(c, "open youtube"), ("task", "open youtube"))
        self.assertEqual(route(c, "now click the first video"), ("task", "now click the first video"))
        c.text_chat.assert_not_called()
    def test_open_url_uses_current_browser(self):
        import intents
        run = mock.MagicMock(); b = mock.MagicMock()
        with mock.patch("actions.IS_MAC", True), mock.patch("intents.browser_dom.frontmost_app", return_value="PC Agent"):
            ex = Executor(mock.MagicMock(), browser=b, runner=run, sleep=lambda s: None)
            ex.execute(parse_action('{"action":"open_app","app":"Safari"}'))
            ex.execute(parse_action('{"action":"open_url","url":"https://www.youtube.com"}'))
        self.assertEqual(run.call_args.args[0], ["open", "-a", "Safari", "https://www.youtube.com"]); b.open.assert_not_called()
        self.assertEqual(intents.STATE["browser"], "Safari")
    def test_frontmost_browser_wins(self):
        import intents
        intents.STATE["browser"] = "Safari"
        with mock.patch("intents.browser_dom.frontmost_app", return_value="Google Chrome"):
            self.assertEqual(intents.current_browser(), "Google Chrome")
    def test_first_video_js_and_fallback(self):
        import intents
        from actions import NeedsVision
        intents.STATE["browser"] = "Safari"
        with mock.patch("actions.IS_MAC", True), mock.patch("intents.browser_dom.frontmost_app", return_value="Safari"):
            with mock.patch("browser_dom.run_js", return_value="ok|Lofi radio") as js:
                ex = Executor(mock.MagicMock(), sleep=lambda s: None); ex.execute({"action": "first_video"})
                self.assertEqual(ex.last_info, "Lofi radio"); self.assertIn("video-title-link", js.call_args.args[1])
            with mock.patch("browser_dom.run_js", side_effect=RuntimeError("JavaScript from Apple Events is turned off")):
                with self.assertRaises(NeedsVision): Executor(mock.MagicMock(), sleep=lambda s: None).execute({"action": "first_video"})

    def _agent(self, client, ex):
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 4, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False, "page_wait_s": 0},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: ex, dom=False)
        a.kill.start = lambda: None; return a
    def test_full_flow_no_model(self):
        client = mock.MagicMock(); ex = mock.MagicMock()
        res = self._agent(client, ex).run("open safari and open youtube and click the first video")
        self.assertTrue(res.startswith("done"))
        self.assertEqual([c.args[0]["action"] for c in ex.execute.call_args_list], ["open_app", "open_url", "first_video"])
        client.chat.assert_not_called(); client.text_chat.assert_not_called()
    def test_first_video_falls_back_to_vision(self):
        from actions import NeedsVision
        client = mock.MagicMock(); client.chat.side_effect = ['not sure', '{"action":"click","x":100,"y":100}', '{"action":"done","summary":"clicked"}']
        ex = mock.MagicMock()
        def execute(a):
            if a["action"] == "first_video": raise NeedsVision("js off")
        ex.execute.side_effect = execute
        self.assertEqual(self._agent(client, ex).run("open youtube and click the first video"), "clicked")
        self.assertIn("click the first video", client.chat.call_args_list[1].args[1])
        self.assertIn("(already done) open_url", client.chat.call_args_list[1].args[1])
    def test_first_video_vision_locate_clicks_points(self):
        from actions import NeedsVision
        client = mock.MagicMock(); client.chat.return_value = '{"x": 250, "y": 500}'
        ex = mock.MagicMock()
        def execute(a):
            if a["action"] == "first_video": raise NeedsVision("js off")
        ex.execute.side_effect = execute
        a = self._agent(client, ex)
        a.capture = lambda: Image.new("RGB", (1600, 1200))  # Retina-like: 2x the 800x600 points
        res = a.run("open youtube and click the first video")
        self.assertTrue(res.startswith("done"), res)
        a.gui.click.assert_called_with(200, 300)  # 250/1000*1280*1.25*0.5 = 200 ; 500/1000*960*1.25*0.5 = 300

class RecentBrowser(unittest.TestCase):
    def test_parse(self):
        import browser_dom
        r = mock.MagicMock(stdout='ASN:0x0-0x1: "PC_Agent", ASN:0x0-0x2: "Finder", ASN:0x0-0x3: "Safari", ASN:0x0-0x4: "Google_Chrome"')
        with mock.patch("browser_dom.sys.platform", "darwin"):
            self.assertEqual(_REAL_RECENT(run=lambda *a, **k: r), "Safari")

class RealRunRegressions(unittest.TestCase):
    def test_string_params(self):
        self.assertEqual(parse_action('{"action":"open_app","params":"Chrome"}')["app"], "Chrome")
    def test_invented_ids_stop_and_no_click_element_offered(self):
        client = mock.MagicMock(); client.chat.return_value = '{"action":"click_element","params":{"id":"video_0"}}'
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 25, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None
        self.assertIn("invalid action", a.run("do the thing"))
        self.assertEqual(client.chat.call_count, 5)
        self.assertNotIn("click_element", client.chat.call_args_list[0].args[0])
        self.assertIn("rejected", client.chat.call_args_list[1].args[1])
    def test_preferred_browser(self):
        import intents
        intents.STATE.update(browser=None, preferred="Safari")
        with mock.patch("intents.browser_dom.frontmost_app", return_value="PC Agent"):
            self.assertEqual(intents.current_browser(), "Safari")
        intents.STATE["preferred"] = None

class BrowserTargeting(unittest.TestCase):
    def setUp(self):
        _intents.STATE.update(browser=None, app=None, preferred=None)
    def test_open_url_uses_last_opened_browser(self):
        cmds = []
        with mock.patch("actions.IS_MAC", True):
            ex = Executor(gui=mock.Mock(), runner=cmds.append, sleep=lambda s: None)
            ex.execute({"action": "open_app", "app": "Safari"})
            ex.execute({"action": "open_url", "url": "youtube.com"})
            ex.execute({"action": "web_search", "query": "cats"})
        self.assertEqual(cmds[1], ["open", "-a", "Safari", "https://youtube.com"])
        self.assertEqual(cmds[2][:3], ["open", "-a", "Safari"])
    def test_frontmost_browser_wins(self):
        _intents.STATE["browser"] = "Safari"
        self.assertEqual(_intents.current_browser(lambda: "Google Chrome"), "Google Chrome")
        self.assertEqual(_intents.current_browser(lambda: "Finder"), "Google Chrome")  # persists across messages
    def test_open_youtube_deterministic(self):
        self.assertEqual(_intents.plan("open youtube")[0], [{"action": "open_url", "url": "https://www.youtube.com"}])
        acts, rest = _intents.plan("open safari and go to youtube then click the first video")
        self.assertEqual([a["action"] for a in acts], ["open_app", "open_url", "first_video"]); self.assertEqual(rest, "")
    def test_first_video_via_js_and_fallback(self):
        with mock.patch("actions.IS_MAC", True), mock.patch("intents.current_browser", return_value="Safari"):
            ex = Executor(gui=mock.Mock(), runner=lambda c: None, sleep=lambda s: None)
            with mock.patch("browser_dom.run_js", return_value="ok|Some video"):
                ex.execute({"action": "first_video"}); self.assertEqual(ex.last_info, "Some video")
            with mock.patch("browser_dom.run_js", side_effect=RuntimeError("not allowed")):
                with self.assertRaises(_actions.NeedsVision): ex.execute({"action": "first_video"})
    def test_imperative_followups_are_tasks(self):
        from core import route
        c = mock.Mock()
        for m in ["click on the first video", "now play the first one", "pause it", "scroll down"]:
            self.assertEqual(route(c, m)[0], "task", m)
        c.text_chat.assert_not_called()

class Round2(unittest.TestCase):
    def test_press_enter_alias(self):
        for raw in ['{"action":"press_enter"}', '{"action":"enter"}', '{"action":"key","key":"enter"}']:
            a = parse_action(raw); self.assertEqual((a["action"], a["keys"]), ("hotkey", ["enter"]))
    def test_search_youtube_deterministic(self):
        acts, rest = _intents.plan("search youtube for 'embrace it' and play the first video")
        self.assertEqual(rest, "")
        self.assertEqual(acts, [{"action": "open_url", "url": "https://www.youtube.com/results?search_query=embrace+it"},
                                {"action": "first_video"}])
        self.assertEqual(_intents.plan("search for embrace it on youtube")[0][0]["url"][-10:], "embrace+it")
    def test_to_points(self):
        from core import to_points
        self.assertEqual(to_points(500, 500, 1280, 831, 3024/1280, 1512/3024), (756, 491))
    def test_retype_is_repeat(self):
        from core import LoopGuard
        g = LoopGuard()
        self.assertEqual(g.check({"action": "type", "text": "embrace it"}), 0)
        g.check({"action": "wait", "seconds": 2})
        self.assertGreaterEqual(g.check({"action": "type", "text": "Embrace it"}), 1)
        self.assertGreaterEqual(g.check({"action": "type", "text": "embrace it"}), 2)
    def test_goal_reached(self):
        self.assertTrue(_intents.goal_reached("play the first video", "https://www.youtube.com/watch?v=x"))
        self.assertFalse(_intents.goal_reached("play the first video", "https://www.youtube.com/results?q"))

class PanelBridge(unittest.TestCase):
    def setUp(self):
        import server, threading, urllib.request
        self.server, self.ur = server, urllib.request
        class FakeAgent:
            def __init__(s, br): s.client = None; s.kill = mock.Mock(); s.kill.event = threading.Event(); s.br = br
            def run(s, task):
                ok = s.br.hub.confirm("click Buy?", timeout=5)
                s.br.log('[1] click {"x": 10, "y": 20}'); return f"ran {task} ok={ok}"
        def router(client, text, hist, on_token=None):
            return ("task", text) if text.startswith("do") else ("chat", "hi " + text)
        br = server.Bridge({}, token="t" * 40, agent_factory=FakeAgent, router=router)
        self.srv, self.br = server.serve({}, port=0, bridge=br) if False else (None, br)
        from http.server import ThreadingHTTPServer
        self.srv = ThreadingHTTPServer(("127.0.0.1", 0), server.make_handler(br)); self.srv.daemon_threads = True
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"
        self.addCleanup(self.srv.shutdown)
    def req(self, path, body=None, token="t" * 40, origin=None):
        import json as j
        h = {"Content-Type": "application/json"}
        if token: h["X-PCA-Token"] = token
        if origin: h["Origin"] = origin
        r = self.ur.Request(self.base + path, data=j.dumps(body).encode() if body is not None else None, headers=h)
        try:
            with self.ur.urlopen(r, timeout=5) as resp: return resp.status, j.loads(resp.read())
        except self.ur.HTTPError as e: return e.code, None
    def test_auth(self):
        self.assertEqual(self.req("/health", token=None)[0], 401)
        self.assertEqual(self.req("/health", token="wrong")[0], 401)
        self.assertEqual(self.req("/health", origin="https://evil.example")[0], 401)  # web pages can't drive it
        self.assertEqual(self.req("/health", origin="chrome-extension://" + "a" * 32)[0], 200)
    def test_task_with_confirm_and_steps(self):
        import time as t
        q = self.br.hub.subscribe()
        self.assertTrue(self.req("/chat", {"text": "do buy it"})[1]["accepted"])
        got, cid = [], None
        end = t.time() + 5
        while t.time() < end:
            m = json.loads(q.get(timeout=5)); got.append(m["type"])
            if m["type"] == "confirm": cid = m["id"]; self.req("/confirm", {"id": cid, "ok": True})
            if m["type"] == "result": self.assertEqual(m["text"], "ran do buy it ok=True"); break
        self.assertIn("step", got); self.assertIn("confirm_done", got); self.assertIsNotNone(cid)
    def test_chat_and_stop(self):
        import time as t
        q = self.br.hub.subscribe(); self.req("/chat", {"text": "hello"})
        while True:
            m = json.loads(q.get(timeout=5))
            if m["type"] == "chat": self.assertEqual(m["text"], "hi hello"); break
        self.assertEqual(self.req("/takeover", {})[0], 200)
    def test_token_file_private(self):
        import tempfile, stat
        p = os.path.join(tempfile.mkdtemp(), "tok")
        t1 = self.server.load_token(p); self.assertGreaterEqual(len(t1), 32); self.assertEqual(self.server.load_token(p), t1)
        if os.name != "nt": self.assertEqual(stat.S_IMODE(os.stat(p).st_mode), 0o600)

class Packaging(unittest.TestCase):
    def test_user_data_dirs(self):
        import paths
        self.assertEqual(paths.user_data_dir("darwin", {}, "/Users/u"), os.path.join("/Users/u", "Library", "Application Support", "Liscia Agent"))
        self.assertTrue(paths.user_data_dir("win32", {"APPDATA": "C:\\\\R"}, "C:\\\\U").endswith("Liscia Agent"))
        self.assertIn("C:\\\\R", paths.user_data_dir("win32", {"APPDATA": "C:\\\\R"}, "C:\\\\U"))
    def test_frozen_config_copied_to_data_dir(self):
        import paths, tempfile
        d = tempfile.mkdtemp()
        with mock.patch.object(paths, "frozen", return_value=True), mock.patch.object(paths, "user_data_dir", return_value=d), mock.patch.object(paths, "resource_dir", return_value=os.path.dirname(os.path.dirname(os.path.abspath(__file__)))):
            p = paths.config_path()
            self.assertEqual(p, os.path.join(d, "config.yaml")); self.assertTrue(os.path.exists(p))
            self.assertEqual(paths.runs_dir("runs"), os.path.join(d, "runs"))
            self.assertEqual(paths.token_path(), os.path.join(d, "token"))
    def test_missing_models(self):
        from llm import missing_models
        self.assertEqual(missing_models(["qwen2.5:7b", "foo:latest"], ["qwen3-vl:4b-instruct", "qwen2.5:7b", "foo", None]),
                         ["qwen3-vl:4b-instruct"])
    def test_pull_progress(self):
        from llm import pull_model
        lines = [b'{"status":"pulling manifest"}', b'{"status":"downloading","completed":50,"total":100}', b'{"status":"success"}']
        resp = mock.MagicMock(); resp.__enter__.return_value = resp; resp.iter_lines.return_value = lines
        seen = []
        self.assertTrue(pull_model("http://h", "m", lambda *a: seen.append(a), post=lambda *a, **k: resp))
        self.assertIn(("downloading", 50, 100), seen)
    def test_setup_models_offers_pull(self):
        import ui, queue, llm
        fake = mock.Mock(); fake.q = queue.Queue(); fake.cfg = {"ollama": {"host": "http://h", "model": "a", "text_model": "b"}}
        with mock.patch.object(llm.OllamaClient, "check", return_value=["b"]), \
             mock.patch("llm.pull_model", side_effect=lambda h, n, p: (p("downloading", 5, 10), True)[1]) as pm:
            self.assertTrue(ui.ChatApp.setup_models(fake, ask=lambda m: True))
            pm.assert_called_once(); self.assertEqual(pm.call_args.args[1], "a")
        msgs = []
        while not fake.q.empty(): msgs.append(fake.q.get())
        self.assertTrue(any(k == "status" and "50%" in v for k, v in msgs))
        with mock.patch.object(llm.OllamaClient, "check", return_value=None), mock.patch("llm.start_ollama", return_value=False):
            self.assertFalse(ui.ChatApp.setup_models(fake, ask=lambda m: True, sleep=lambda s: None))

class Rename(unittest.TestCase):
    def test_migrates_pc_agent_data(self):
        import paths, tempfile
        root = tempfile.mkdtemp(); old = os.path.join(root, "PC Agent"); os.makedirs(old)
        open(os.path.join(old, "token"), "w").write("x" * 40); open(os.path.join(old, "config.yaml"), "w").write("a: 1\n")
        new = os.path.join(root, "Liscia Agent")
        with mock.patch.object(paths, "frozen", return_value=True), mock.patch.object(paths, "user_data_dir", return_value=new):
            self.assertEqual(paths.data_dir(), new)
        self.assertEqual(open(os.path.join(new, "token")).read(), "x" * 40)
        self.assertTrue(os.path.exists(os.path.join(new, "config.yaml")))
        self.assertFalse(paths.migrate_old_data(new, old))  # only once
    def test_extension_manifest(self):
        m = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "chrome-extension", "manifest.json")))
        self.assertIn("Liscia", m["name"]); self.assertEqual(set(m["icons"]), {"16", "32", "48", "128"})

class FakeExt(threading.Thread if False else object):
    """A mocked extension: polls an ExtensionLink and answers commands with a tiny fake page."""
    def __init__(self, link):
        import threading
        self.link, self.calls, self.stop = link, [], False
        self.url, self.typed = "https://www.youtube.com/", ""
        self.t = threading.Thread(target=self.loop, daemon=True); self.t.start()
    def snap(self):
        els = [{"id": "1", "role": "searchbox", "label": "Search", "value": self.typed}, {"id": "2", "role": "button", "label": "Search"}]
        if "results" in self.url: els.append({"id": "7", "role": "link", "label": "Embrace It (Official Video)", "href": "https://www.youtube.com/watch?v=x"})
        return {"url": self.url, "title": "Embrace It (Official Video) - YouTube" if "watch" in self.url else "YouTube", "elements": els, "dialogs": [], "text": "YouTube", "scroll": [0, 100]}
    def loop(self):
        while not self.stop:
            for c in self.link.poll({"version": "test"}, hold=0.05):
                self.calls.append(c["cmd"]); a = c["args"]
                if c["cmd"] == "snapshot": r = {"id": c["id"], "ok": True, "result": self.snap()}
                elif c["cmd"] == "type": self.typed = a["text"]; self.url = "https://www.youtube.com/results?search_query=" + a["text"].replace(" ", "+") if a.get("press_enter") else self.url; r = {"id": c["id"], "ok": True, "result": {}}
                elif c["cmd"] == "click" and a["id"] == "7": self.url = "https://www.youtube.com/watch?v=x"; r = {"id": c["id"], "ok": True, "result": {"clicked": "Embrace It"}}
                elif c["cmd"] == "navigate": self.url = a["url"]; r = {"id": c["id"], "ok": True, "result": {"title": "YouTube"}}
                elif c["cmd"] == "first_video": self.url = "https://www.youtube.com/watch?v=x"; r = {"id": c["id"], "ok": True, "result": {"clicked": "Embrace It"}}
                else: r = {"id": c["id"], "ok": False, "error": "nope"}
                self.link.result(r)

import threading
class ExtensionChannel(unittest.TestCase):
    def setUp(self):
        import extlink
        self.link = extlink.ExtensionLink(); self.ext = FakeExt(self.link); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            import time; time.sleep(0.01)
    def test_roundtrip_and_status(self):
        self.assertTrue(self.link.connected())
        self.assertEqual(self.link.call("snapshot", timeout=3)["url"], "https://www.youtube.com/")
        import extlink
        with self.assertRaises(extlink.ExtError): self.link.call("bogus", timeout=3)
        st = self.link.status(); self.assertTrue(st["connected"]); self.assertEqual(st["extension"]["version"], "test")
    def test_not_connected(self):
        import extlink
        l = extlink.ExtensionLink()
        with self.assertRaises(extlink.ExtError): l.call("snapshot", timeout=0.1)
    def test_web_loop_search_and_click(self):
        import webagent
        client = mock.MagicMock(); client.text_chat.side_effect = [
            '{"thought":"type query","action":"type","id":"1","text":"embrace it","press_enter":true}',
            '{"thought":"first result","action":"click","id":"7"}']
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False
        ag.cfg = {"safety": {"confirm_risky": True, "risky_words": ["buy"]}}; logs = []; ag.log = logs.append
        res = webagent.run_web(ag, "search youtube for embrace it and play the first video", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None)
        self.assertIn("Playing \"Embrace It (Official Video)\"", res)
        self.assertEqual([c for c in self.ext.calls if c != "snapshot"], ["type", "click"])
        self.assertTrue(any("Embrace It" in l for l in logs))
        ag.confirm.assert_not_called()
    def test_web_loop_risky_needs_confirm(self):
        import webagent
        self.ext.snap = lambda: {"url": "https://shop/", "title": "Shop", "elements": [{"id": "3", "role": "button", "label": "Buy now"}], "dialogs": []}
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"click","id":"3"}', '{"action":"done","summary":"stopped"}']
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.confirm.return_value = False
        ag.cfg = {"safety": {"confirm_risky": True, "risky_words": ["buy"]}}
        self.assertEqual(webagent.run_web(ag, "buy it", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None), "stopped")
        ag.confirm.assert_called_once(); self.assertNotIn("click", self.ext.calls)
    def test_parse_web_actions(self):
        import webagent
        self.assertEqual(webagent.parse('{"action":"click_element","id":"[12]"}'), {"action": "click", "id": "12"})
        self.assertEqual(webagent.parse('{"action":"press_enter"}')["key"], "Enter")
        with self.assertRaises(ActionError): webagent.parse('{"action":"click"}')
    def test_open_url_and_first_video_via_extension(self):
        import extlink
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("actions.IS_MAC", True), \
             mock.patch("intents.current_browser", return_value="Google Chrome"):
            cmds = []; ex = Executor(gui=mock.Mock(), runner=cmds.append, sleep=lambda s: None)
            ex.execute({"action": "open_url", "url": "https://www.youtube.com/results?search_query=x"})
            ex.execute({"action": "first_video"})
        self.assertEqual(cmds, [["open", "-a", "Google Chrome"]])
        self.assertIn("navigate", self.ext.calls); self.assertIn("first_video", self.ext.calls); self.assertEqual(ex.last_info, "Embrace It")
    def test_open_app_chrome_alias(self):
        cmds = []
        with mock.patch("actions.IS_MAC", True):
            Executor(gui=mock.Mock(), runner=cmds.append, sleep=lambda s: None).execute({"action": "open_app", "app": "Chrome"})
        self.assertEqual(cmds[0], ["open", "-a", "Google Chrome"])
    def test_http_endpoints(self):
        import server, json as j, urllib.request as ur
        from http.server import ThreadingHTTPServer
        br = server.Bridge({}, token="t" * 40)
        br.link = self.link
        srv = ThreadingHTTPServer(("127.0.0.1", 0), server.make_handler(br)); threading.Thread(target=srv.serve_forever, daemon=True).start()
        self.addCleanup(srv.shutdown)
        req = ur.Request(f"http://127.0.0.1:{srv.server_address[1]}/status", headers={"X-PCA-Token": "t" * 40})
        st = j.loads(ur.urlopen(req, timeout=5).read())
        self.assertTrue(st["extension"]["connected"])
        req = ur.Request(f"http://127.0.0.1:{srv.server_address[1]}/ext/poll?hold=0.1", headers={"X-PCA-Token": "bad" * 14})
        with self.assertRaises(ur.HTTPError): ur.urlopen(req, timeout=5)

class AgentUsesExtension(ExtensionChannel):
    def test_agent_routes_browser_task_to_extension(self):
        import extlink
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"click","id":"2"}', '{"action":"done","summary":"clicked search"}']
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 4, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link
        with mock.patch("intents.current_browser", return_value="Google Chrome"), mock.patch.object(extlink, "LINK", self.link):
            self.assertEqual(a.run("press the search button on the page"), "clicked search")
        client.chat.assert_not_called()  # no screenshots

class V121WebRouting(ExtensionChannel):
    """v1.2.1: web tasks go through the extension in Chrome; never Safari unless named; no vision hallucination loop."""
    def _agent(self, client, cap=None):
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 6, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        self.execs = []
        def mk(s):
            ex = mock.MagicMock(); ex.execute.side_effect = self.execs.append; return ex
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=cap or (lambda: Image.new("RGB", (800, 600))), executor=mk, dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None
        return a
    def test_open_youtube_in_chrome_and_search(self):
        import extlink
        client = mock.MagicMock()
        a = self._agent(client)
        with mock.patch.object(extlink, "LINK", self.link):
            res = a.run("open youtube in chrome and search lofi music")
        self.assertTrue(res.startswith('Searched YouTube for "lofi music"'), res)
        self.assertEqual(self.ext.url, "https://www.youtube.com/results?search_query=lofi+music")
        client.chat.assert_not_called(); client.text_chat.assert_not_called(); self.assertEqual(self.execs, [])
    def test_search_youtube_for(self):
        import extlink
        client = mock.MagicMock(); a = self._agent(client)
        with mock.patch.object(extlink, "LINK", self.link):
            res = a.run("search youtube for lofi music")
        self.assertTrue(res.startswith('Searched YouTube for "lofi music"'), res); self.assertIn("results?search_query=lofi+music", self.ext.url)
        client.chat.assert_not_called()
    def test_free_form_web_task_uses_webagent_not_vision(self):
        import extlink
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"type","id":"1","text":"daft punk","press_enter":true}', '{"action":"done","summary":"searched"}']
        a = self._agent(client); a.cfg["agent"]["planner"] = False
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("intents.current_browser", return_value="Safari"):
            self.assertEqual(a.run("go on the page and look for daft punk songs"), "searched")
        client.chat.assert_not_called(); self.assertIn("type", self.ext.calls)
    def test_named_other_browser_skips_extension(self):
        import extlink
        a = self._agent(mock.MagicMock())
        with mock.patch.object(extlink, "LINK", self.link):
            self.assertFalse(a.use_extension("open youtube in safari"))
            self.assertTrue(a.use_extension("open youtube in chrome"))
            self.assertTrue(a.use_extension("search google for weather"))
            self.assertFalse(a.use_extension("open notes"))
    def test_never_switch_to_safari_unless_named(self):
        a = self._agent(mock.MagicMock())
        _intents.STATE["named"] = None
        self.assertEqual(a.fix_browser({"action": "open_app", "app": "Safari"})["app"], "Google Chrome")
        self.assertEqual(a.fix_browser({"action": "open_url", "url": "u", "browser": "Safari"})["browser"], "Google Chrome")
        self.assertEqual(a.fix_browser({"action": "open_app", "app": "Notes"})["app"], "Notes")
        _intents.STATE["named"] = "Safari"
        self.assertEqual(a.fix_browser({"action": "open_app", "app": "Safari"})["app"], "Safari")
        _intents.STATE["named"] = None
    def test_current_browser_prefers_chrome_when_extension_live(self):
        import extlink
        _intents.STATE.update(named=None, browser="Safari")
        with mock.patch.object(extlink, "LINK", self.link):
            self.assertEqual(_intents.current_browser(frontmost=lambda: "Safari"), "Google Chrome")
        _intents.STATE["browser"] = None
    def test_vision_fallback_told_chrome_is_open_and_stops_open_loop(self):
        import extlink
        self.ext.loop_tabs = True
        orig = self.ext.loop
        client = mock.MagicMock()
        client.chat.side_effect = ['{"thought":"Chrome is not open","action":"open_app","app":"Safari"}',
                                   '{"thought":"Chrome is not open","action":"open_app","app":"Chrome"}',
                                   '{"action":"open_url","url":"https://x.test/"}', '{"action":"open_url","url":"https://x.test/"}',
                                   '{"action":"open_url","url":"https://x.test"}', '{"action":"done","summary":"no"}']
        a = self._agent(client)
        a.cfg["agent"]["max_steps"] = 6
        with mock.patch.object(extlink, "LINK", self.link), mock.patch.object(a, "live_tab", return_value=("https://www.youtube.com/", "YouTube")), \
             mock.patch.object(a, "use_extension", return_value=False):
            res = a.run("do the thing on screen")
        self.assertIn("kept re-opening", res)
        prompt = client.chat.call_args_list[0][0][1]
        self.assertIn("Google Chrome IS open", prompt)
        self.assertEqual(self.execs[0]["app"], "Google Chrome")  # Safari rewritten
        self.assertFalse(any(e.get("app") == "Safari" for e in self.execs))
    def test_plan_contextual_youtube_search(self):
        acts, rest = _intents.plan("open youtube in chrome and search never gonna give you up")
        self.assertEqual(rest, ""); self.assertIn("search_query=never+gonna+give+you+up", acts[-1]["url"])
        acts, rest = _intents.plan("search youtube for lofi in chrome")
        self.assertEqual(acts[0]["url"], "https://www.youtube.com/results?search_query=lofi")
        self.assertEqual(_intents.named_browser("open youtube in Chrome"), "Google Chrome")
        self.assertIsNone(_intents.named_browser("open youtube"))

class ScrollFakeExt(FakeExt):
    """YouTube results where the wanted video only appears after 3 scrolls (lazy-loaded)."""
    def __init__(self, link):
        self.scrolled = 0; super().__init__(link)
    def snap(self):
        els = [{"id": "1", "role": "searchbox", "label": "Search", "value": ""}, {"id": "2", "role": "button", "label": "Search"},
               {"id": "5", "role": "link", "label": "Pursuit (Live)", "href": "https://www.youtube.com/watch?v=a"}]
        if self.scrolled >= 3:
            els.append({"id": "9", "role": "link", "label": "Pursuit of Jade OST", "href": "https://www.youtube.com/watch?v=jade", "below": True})
            els.append({"id": "10", "role": "link", "label": "Pursuit - isq (Official Audio)", "href": "https://www.youtube.com/watch?v=isq", "below": True})
        titles = {"jade": "Pursuit of Jade OST", "isq": "Pursuit - isq (Official Audio)", "a": "Pursuit (Live)"}
        cur = self.url.split("watch?v=")[-1] if "watch?v=" in self.url else ""
        return {"url": self.url, "title": (titles.get(cur, "") + " - YouTube") if cur else "YouTube", "elements": els, "dialogs": [], "text": "", "scroll": [self.scrolled * 800, 8000]}
    def loop(self):
        while not self.stop:
            for c in self.link.poll({"version": "test"}, hold=0.05):
                self.calls.append(c["cmd"]); a = c["args"]; r = {"id": c["id"], "ok": True, "result": {}}
                if c["cmd"] == "snapshot": r["result"] = self.snap()
                elif c["cmd"] == "navigate": self.url = a["url"]
                elif c["cmd"] == "scroll": self.scrolled += 1; r["result"] = {"moved": 800, "grew": True, "at_bottom": False}
                elif c["cmd"] == "find_text":
                    ok = self.scrolled >= 3 and "jade" in a["text"].lower()
                    r["result"] = {"found": True, "id": "9", "label": "Pursuit of Jade OST"} if ok else {"found": False}
                elif c["cmd"] == "click":
                    if a["id"] == "9": self.url = "https://www.youtube.com/watch?v=jade"
                elif c["cmd"] == "first_video": self.url = "https://www.youtube.com/watch?v=a"
                self.link.result(r)

class V122Scrolling(unittest.TestCase):
    def setUp(self):
        import extlink, time
        self.link = extlink.ExtensionLink(); self.ext = ScrollFakeExt(self.link); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            time.sleep(0.01)
    def _ag(self, client):
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.log = lambda m: None
        ag.cfg = {"safety": {"confirm_risky": False, "risky_words": []}}; return ag
    def test_scroll_find_scrolls_until_found(self):
        import webagent
        res = webagent.scroll_find(self.link, "pursuit of jade", log=lambda m: None, sleep=lambda s: None)
        self.assertEqual(res["id"], "9"); self.assertEqual(self.ext.calls.count("scroll"), 3)
    def test_scroll_find_gives_up_after_max(self):
        import webagent
        self.assertIsNone(webagent.scroll_find(self.link, "nonexistent thing", log=lambda m: None, max_scrolls=5, sleep=lambda s: None))
        self.assertEqual(self.ext.calls.count("scroll"), 5)
    def test_done_not_found_is_replaced_by_scrolling(self):
        import webagent
        client = mock.MagicMock(); client.text_chat.side_effect = [
            '{"action":"done","summary":"Pursuit of Jade not found"}', '{"action":"done","summary":"could not find it"}',
            '{"action":"done","summary":"not found"}', '{"action":"click","id":"9"}', '{"action":"done","summary":"playing"}']
        res = webagent.run_web(self._ag(client), "play pursuit of jade", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None)
        self.assertIn("Pursuit of Jade OST", res); self.assertEqual(self.ext.url, "https://www.youtube.com/watch?v=jade"); self.assertEqual(self.ext.calls.count("scroll"), 3); self.assertIn("click", self.ext.calls)
    def test_find_action(self):
        import webagent
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"find","text":"pursuit of jade"}', '{"action":"click","id":"9"}', '{"action":"done","summary":"ok"}']
        self.assertEqual(webagent.run_web(self._ag(client), "open pursuit of jade", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None), "ok")
        self.assertEqual(self.ext.url, "https://www.youtube.com/watch?v=jade")
        self.assertIn("(below)", webagent.format_snapshot(self.ext.snap()))
    def test_deterministic_play_title_on_youtube_scrolls(self):
        import extlink
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 4, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        client = mock.MagicMock()
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"):
            res = a.run("play pursuit of jade on youtube")
        self.assertIn('Playing "Pursuit of Jade OST"', res); self.assertEqual(self.ext.url, "https://www.youtube.com/watch?v=jade")
        client.text_chat.assert_not_called(); client.chat.assert_not_called()
    def test_plan_play_title(self):
        acts, rest = _intents.plan("play pursuit of jade ost on youtube")
        self.assertEqual(rest, ""); self.assertEqual(acts[1], {"action": "first_video", "match": "pursuit of jade ost"})
        acts, rest = _intents.plan("open youtube and search pursuit and click pursuit of jade")
        self.assertEqual(rest, ""); self.assertEqual(acts[-1]["match"], "pursuit of jade")
        self.assertEqual(_intents.plan("search youtube for x and play the first video")[0][-1], {"action": "first_video"})

USER_SENTENCE = "can you open up youtube and on youtube can you search for the song Pursuit by isq and claick on to play the song"
class V123Tolerant(unittest.TestCase):
    def test_user_exact_sentence(self):
        self.assertEqual(_intents.youtube_request(USER_SENTENCE), ("Pursuit by isq", True))
        acts, rest = _intents.plan(USER_SENTENCE)
        self.assertEqual(rest, "")
        self.assertEqual(acts, [{"action": "open_url", "url": "https://www.youtube.com/results?search_query=Pursuit+by+isq"},
                                {"action": "first_video", "match": "pursuit isq"}])
        self.assertTrue(_intents.is_web_task(USER_SENTENCE))
    def test_variants(self):
        cases = {
            "open youtube and search pursuit by isq": ("pursuit by isq", False),
            "Open up YouTube, search for Pursuit by isq and click on it": ("Pursuit by isq", True),
            "on youtube serach the song pursuit by isq then play it": ("pursuit by isq", True),
            "can you search youtube for pursuit by isq and play the first video": ("pursuit by isq", True),
            "play pursuit by isq on youtube": ("pursuit by isq", True),
            "please go to youtube and look up lofi beats": ("lofi beats", False),
            "search for 'Pursuit' by isq on youtube and clik to play": ("Pursuit' by isq", True) if False else None,
            "open yotube and search for the video called never gonna give you up and watch it": ("never gonna give you up", True),
            "search youtube for lofi in chrome": ("lofi", False),
        }
        for text, want in cases.items():
            if want is None: continue
            self.assertEqual(_intents.youtube_request(text), want, text)
        self.assertIsNone(_intents.youtube_request("open notes and type hello"))
    def test_web_routing_words(self):
        for t in ("claick the login button", "play some music", "go to the website", "serach for cats"):
            self.assertTrue(_intents.is_web_task(t), t)
        self.assertFalse(_intents.is_web_task("open notes"))
    def test_blank_screenshot(self):
        self.assertTrue(screen.is_blank(Image.new("RGB", (800, 600))))
        img = Image.new("RGB", (800, 600)); img.paste((255, 255, 255), (0, 0, 400, 300))
        self.assertFalse(screen.is_blank(img))

class V123Agent(V122Scrolling):
    def _agent(self, client, cap=None):
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 4, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=cap or (lambda: Image.new("RGB", (800, 600))), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None; return a
    def test_user_sentence_runs_on_extension_without_models(self):
        import extlink
        self.ext.find = None
        orig = self.ext.loop
        client = mock.MagicMock(); a = self._agent(client)
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"):
            res = a.run(USER_SENTENCE)
        self.assertIn('Playing "Pursuit - isq', res)  # title AND artist matched, not "Pursuit (Live)" or "Pursuit of Jade"
        self.assertIn("navigate", self.ext.calls); self.assertEqual(self.ext.url, "https://www.youtube.com/watch?v=isq")
        client.chat.assert_not_called(); client.text_chat.assert_not_called()
    def test_vision_stops_on_blank_screenshots(self):
        client = mock.MagicMock(); a = self._agent(client)
        a.cfg["agent"]["blank_check"] = True
        with mock.patch.object(a, "use_extension", return_value=False):
            res = a.run("do something on the desktop")
        self.assertIn("Screen Recording", res); client.chat.assert_not_called()
    def test_extension_error_does_not_fall_into_vision(self):
        import extlink, webagent
        client = mock.MagicMock(); a = self._agent(client)
        with mock.patch.object(extlink, "LINK", self.link), mock.patch.object(webagent, "run_web", side_effect=RuntimeError("boom")):
            res = a.run("click the login button on the page")
        self.assertIn("extension step failed", res); client.chat.assert_not_called()
    def test_restricted_tab_opens_site_first(self):
        import webagent, extlink
        calls = {"n": 0}; real = self.link.call
        def call(cmd, **kw):
            if cmd == "snapshot" and calls["n"] == 0:
                calls["n"] += 1; raise extlink.ExtError("can't access chrome://extensions/ (browser page)")
            return real(cmd, **kw)
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"done","summary":"ok"}']
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.log = lambda m: None
        ag.cfg = {"safety": {"confirm_risky": False, "risky_words": []}}
        with mock.patch.object(self.link, "call", side_effect=call):
            self.assertEqual(webagent.run_web(ag, "on youtube click something", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None), "ok")
        self.assertEqual(self.ext.url, "https://www.youtube.com")

class V124Slash(unittest.TestCase):
    def test_token_shown_and_copied(self):
        import slash
        got = []
        r = slash.handle("/token", copy=got.append, token_loader=lambda: "T" * 43)
        self.assertIn("T" * 43, r); self.assertIn("Chrome panel", r); self.assertIn("Copied", r); self.assertEqual(got, ["T" * 43])
        r = slash.handle(" /TOKEN ", copy=None, token_loader=lambda: "abc" * 15); self.assertNotIn("Copied", r)
    def test_copy_failure_is_ok(self):
        import slash
        def bad(t): raise RuntimeError("no clipboard")
        self.assertNotIn("Copied", slash.handle("/token", copy=bad, token_loader=lambda: "x" * 40))
    def test_help_and_passthrough(self):
        import slash
        self.assertIn("/token", slash.handle("/help")); self.assertIsNone(slash.handle("open youtube"))
        self.assertIsNone(slash.handle("/do open notes")); self.assertIsNone(slash.handle(""))
    def test_ui_never_sends_token_to_model(self):
        import ui, slash
        app = ui.ChatApp.__new__(ui.ChatApp)
        app.busy = False; app.entry = mock.MagicMock(); app.entry.get.return_value = "/token"
        writes = []; app.write = lambda tag, t: writes.append((tag, t)); app._copy = mock.MagicMock(); app._set_busy = mock.MagicMock(); app.q = mock.MagicMock()
        with mock.patch("server.load_token", return_value="S" * 43), mock.patch("threading.Thread") as th:
            app.send()
        th.assert_not_called(); app._set_busy.assert_not_called()
        self.assertTrue(any("S" * 43 in t for _, t in writes)); app._copy.assert_called_once_with("S" * 43)
    def test_bridge_slash_not_routed(self):
        import server
        router = mock.MagicMock(); br = server.Bridge({}, token="t" * 40, router=router)
        ev = []; br.hub.emit = lambda kind, **d: ev.append((kind, d))
        br._chat("/token"); br._chat("/help")
        router.assert_not_called()
        self.assertFalse(any("t" * 40 in str(d) for _, d in ev)); self.assertTrue(any("/token" in str(d) for _, d in ev))

WIN_TASK = "can you go to google and search for who the top 10 richest people are in the world right now"
class GoogleFake(FakeExt):
    def __init__(self, link, results_text="", page_text=""):
        self.results_text, self.page_text = results_text, page_text; super().__init__(link)
    def snap(self):
        els = [{"id": "1", "role": "link", "label": "Privacy", "href": "https://policies.google.com/privacy"},
               {"id": "2", "role": "link", "label": "Jobs", "href": "https://careers.google.com"},
               {"id": "3", "role": "link", "label": "The World's Real-Time Billionaires - Forbes", "href": "https://www.forbes.com/real-time-billionaires/"}]
        title = "Not Found" if "top-rich-list" in self.url else "results"
        return {"url": self.url, "title": title, "elements": els, "dialogs": [], "text": "404 page not found" if title == "Not Found" else "", "scroll": [0, 100]}
    def loop(self):
        while not self.stop:
            for c in self.link.poll({"version": "test"}, hold=0.05):
                self.calls.append((c["cmd"], c["args"])); a = c["args"]; r = {"id": c["id"], "ok": True, "result": {}}
                if c["cmd"] == "snapshot": r["result"] = self.snap()
                elif c["cmd"] == "navigate": self.url = a["url"]
                elif c["cmd"] == "extract_text": r["result"] = {"text": self.page_text if "forbes" in self.url else self.results_text}
                self.link.result(r)

class V125Answers(unittest.TestCase):
    def setUp(self):
        import extlink, time
        self.link = extlink.ExtensionLink()
    def start(self, **kw):
        import time
        self.ext = GoogleFake(self.link, **kw); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            time.sleep(0.01)
    def agent(self, client):
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 40, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None; return a
    def test_parse_google_requests(self):
        self.assertEqual(_intents.google_request(WIN_TASK), "who the top 10 richest people are in the world right now")
        self.assertEqual(_intents.google_request("google best pizza near me"), "best pizza near me")
        self.assertEqual(_intents.google_request("search the web for python 3.14 release date"), "python 3.14 release date")
        self.assertEqual(_intents.google_request("search for cheap flights to tokyo on google"), "cheap flights to tokyo")
        self.assertIsNone(_intents.google_request("search youtube for lofi"))
        self.assertTrue(_intents.is_question(WIN_TASK)); self.assertFalse(_intents.is_question("search for cat pictures"))
        acts, rest = _intents.plan(WIN_TASK)
        self.assertEqual(rest, ""); self.assertIn("google.com/search?q=who+the+top+10+richest", acts[0]["url"]); self.assertNotIn("2023", acts[0]["url"])
        self.assertEqual(acts[1]["action"], "answer_from_page")
    def test_answers_from_results_page_in_one_model_call(self):
        import extlink
        self.start(results_text="1. Elon Musk $400B 2. Larry Ellison ...")
        client = mock.MagicMock(); client.text_chat.return_value = '{"answer": "1. Elon Musk\\n2. Larry Ellison"}'
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"):
            res = self.agent(client).run(WIN_TASK)
        self.assertEqual(res, "1. Elon Musk\n2. Larry Ellison"); self.assertEqual(client.text_chat.call_count, 1)
        sys_prompt = client.text_chat.call_args[0][0][0]["content"]
        import datetime; self.assertIn(str(datetime.date.today().year), sys_prompt)
    def test_opens_top_result_when_results_page_insufficient(self):
        import extlink
        self.start(results_text="ads", page_text="Forbes real-time list: Musk, Ellison, Zuckerberg")
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"answer": "", "need_more": true}', '{"answer": "Musk, Ellison, Zuckerberg"}']
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"):
            res = self.agent(client).run(WIN_TASK)
        self.assertEqual(res, "Musk, Ellison, Zuckerberg")
        navs = [a["url"] for c, a in self.ext.calls if c == "navigate"]
        self.assertEqual(navs[-1], "https://www.forbes.com/real-time-billionaires/")  # not Privacy/Jobs
    def test_webagent_blocks_repeats_404_and_junk_and_answers(self):
        import webagent
        self.start(results_text="Musk is first")
        self.ext.url = "https://www.google.com/search?q=x"
        client = mock.MagicMock(); client.text_chat.side_effect = [
            '{"action":"click","id":"1"}',                                              # Privacy -> blocked
            '{"action":"navigate","url":"https://www.forbes.com/top-rich-list"}',       # 404
            '{"action":"navigate","url":"https://www.forbes.com/top-rich-list"}',       # blocked (dead url)
            '{"action":"navigate","url":"https://www.google.com/search?q=x"}',
            '{"action":"navigate","url":"https://www.google.com/search?q=x"}',
            '{"action":"navigate","url":"https://www.google.com/search?q=x"}',          # 3rd identical -> blocked
            '{"action":"read_page"}', '{"action":"answer","text":"Elon Musk is the richest."}']
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.log = lambda m: None
        ag.cfg = {"safety": {"confirm_risky": False, "risky_words": []}, "agent": {}}
        res = webagent.run_web(ag, "who is the richest person", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None)
        self.assertEqual(res, "Elon Musk is the richest.")
        navs = [a["url"] for c, a in self.ext.calls if c == "navigate"]
        self.assertEqual(navs.count("https://www.forbes.com/top-rich-list"), 0)  # made-up deep link -> Google site search
        self.assertTrue(any("site%3Aforbes.com" in n for n in navs), navs)
        self.assertEqual(navs.count("https://www.google.com/search?q=x"), 2)
        self.assertNotIn("click", [c for c, a in self.ext.calls])
        prompt = client.text_chat.call_args_list[0][0][0][1]["content"]
        self.assertIn("TODAY:", prompt); self.assertNotIn('"Privacy"', prompt); self.assertNotIn('"Jobs"', prompt)
    def test_info_question_step_cap(self):
        import webagent
        self.start()
        client = mock.MagicMock(); client.text_chat.return_value = '{"action":"scroll","direction":"down"}'
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.log = lambda m: None
        ag.cfg = {"safety": {"confirm_risky": False, "risky_words": []}, "agent": {}}
        webagent.run_web(ag, "what is the weather in paris", [], open(os.devnull, "w"), link=self.link, max_steps=40, sleep=lambda s: None)
        self.assertLessEqual(client.text_chat.call_count, 8)
    def test_parse_answer_actions(self):
        import webagent
        self.assertEqual(webagent.parse('{"action":"extract_text"}')["action"], "read_page")
        self.assertEqual(webagent.parse('{"action":"answer","answer":"x"}')["text"], "x")
        with self.assertRaises(ActionError): webagent.parse('{"action":"answer"}')

class ExtensionZip(unittest.TestCase):
    def test_zip_stamped_with_version(self):
        import tempfile, zipfile, importlib.util
        root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        spec = importlib.util.spec_from_file_location("bez", os.path.join(root, "scripts", "build_extension_zip.py"))
        m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
        with tempfile.TemporaryDirectory() as d:
            out = m.build("v9.8.7", out_dir=d)
            self.assertTrue(out.endswith("liscia-agent-chrome-extension-9.8.7.zip"))
            with zipfile.ZipFile(out) as z:
                names = z.namelist(); man = json.loads(z.read("manifest.json"))
            self.assertEqual(man["version"], "9.8.7")
            for f in ("background.js", "content.js", "sidepanel.html", "icon128.png"): self.assertIn(f, names)
            with self.assertRaises(SystemExit): m.build("1.2.5-beta", out_dir=d)

class V126OpenSite(unittest.TestCase):
    def test_plan_open_site_in_browser(self):
        _intents.guess_site.cache_clear()
        for t in ("open Hulu in the browser", "open hulu in my browser", "go to hulu", "open hulu in chrome", "can you open up hulu in the web browser"):
            acts, rest = _intents.plan(t)
            self.assertEqual((acts[0]["url"], rest), ("https://www.hulu.com", ""), t); self.assertEqual(len(acts), 1)
            self.assertTrue(_intents.is_plain_open(t), t)
        acts, _ = _intents.plan("go to crunchyroll"); self.assertEqual(acts[0]["url"], "https://www.crunchyroll.com")
        self.assertEqual(_intents.plan("open notes")[0][0]["action"], "open_app")
        self.assertEqual(_intents.plan("open calculator")[0], [])  # no browser cue -> no .com guess
        self.assertFalse(_intents.is_plain_open("open youtube and search lofi"))
    def test_guess_falls_back_to_lucky_google(self):
        _intents.guess_site.cache_clear()
        with mock.patch.object(_intents, "GUESS_CHECK", True), mock.patch.object(_intents, "site_alive", return_value=False):
            self.assertIn("google.com/search?btnI=1&q=zzqqx", _intents.guess_site("zzqqx"))
        _intents.guess_site.cache_clear()
    def test_404_only_from_status_or_title(self):
        import webagent
        self.assertFalse(webagent.is_404({"title": "Sign in | Hulu", "text": "error 404 page not found somewhere in text"}))
        self.assertTrue(webagent.is_404({"title": "Hulu", "status": 404}))
        self.assertFalse(webagent.is_404({"title": "404 - fun facts", "status": 200}))
        self.assertTrue(webagent.is_404({"title": "Page Not Found"})); self.assertTrue(webagent.is_404({"title": "Error 404 | Forbes"}))
    def test_money_clicks(self):
        import webagent
        el = {"id": "5", "role": "link", "label": "Sign up for only Hulu", "href": "https://signup.hulu.com/plans"}
        self.assertTrue(webagent.money_click({"action": "click", "id": "5"}, el, "open hulu"))
        self.assertFalse(webagent.money_click({"action": "click", "id": "5"}, el, "sign up for hulu"))
        self.assertFalse(webagent.money_click({"action": "click", "id": "5"}, {"label": "Log In"}, "open hulu"))
        self.assertTrue(webagent.site_matches("https://www.hulu.com/welcome", "https://www.hulu.com"))
        self.assertFalse(webagent.site_matches("https://www.google.com/x", "https://www.hulu.com"))

class HuluFake(FakeExt):
    def snap(self):
        return {"url": self.url, "title": "Hulu", "elements": [{"id": "5", "role": "link", "label": "Sign up for only Hulu", "href": "https://signup.hulu.com/plans"}],
                "dialogs": [], "text": "", "scroll": [0, 100], "status": 200}
    def loop(self):
        while not self.stop:
            for c in self.link.poll({"version": "test"}, hold=0.05):
                self.calls.append(c["cmd"]); a = c["args"]; r = {"id": c["id"], "ok": True, "result": {}}
                if c["cmd"] == "snapshot": r["result"] = self.snap()
                elif c["cmd"] == "navigate": self.url = a["url"] if a["url"] != "https://www.hulu.com" else "https://www.hulu.com/welcome"
                self.link.result(r)

class V126Agent(unittest.TestCase):
    def setUp(self):
        import extlink, time
        self.link = extlink.ExtensionLink(); self.ext = HuluFake(self.link); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            time.sleep(0.01)
    def test_open_hulu_one_step_and_stop(self):
        import extlink
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 10, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": True, "kill_keys": []}}
        client = mock.MagicMock()
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"):
            res = a.run("open Hulu in the browser")
        self.assertTrue(res.startswith("done"), res)
        self.assertEqual([c for c in self.ext.calls if c != "snapshot"], ["navigate"])
        client.text_chat.assert_not_called(); client.chat.assert_not_called()
    def test_webagent_asks_before_signup_and_stops_after_open(self):
        import webagent
        client = mock.MagicMock(); client.text_chat.side_effect = ['{"action":"click","id":"5"}', '{"action":"navigate","url":"https://www.hulu.com"}']
        ag = mock.MagicMock(); ag.client = client; ag.kill.killed = False; ag.log = lambda m: None; ag.confirm.return_value = False
        ag.cfg = {"safety": {"confirm_risky": True, "risky_words": []}, "agent": {}}
        res = webagent.run_web(ag, "open Hulu in the browser", [], open(os.devnull, "w"), link=self.link, sleep=lambda s: None)
        self.assertTrue(res.startswith("done: opened"), res); ag.confirm.assert_called_once()
        self.assertNotIn("click", self.ext.calls)


DOC_TASK = "can you open up forbes and looks for the top 20 richest people and summrize and put it in a word document"
class V127Docs(unittest.TestCase):
    def test_parse_mixed_task(self):
        o = _intents.output_request(DOC_TASK)
        self.assertEqual((o["kind"], o["research"]), ("docx", "can you open up forbes and looks for the top 20 richest people"))
        self.assertEqual(_intents.research_query(o["research"]), ("forbes", "forbes top 20 richest people"))
        self.assertEqual(_intents.requested_count(DOC_TASK), 20); self.assertTrue(_intents.is_web_task(DOC_TASK))
        self.assertEqual(_intents.output_request("find laptop deals and save them in an excel spreadsheet")["kind"], "xlsx")
        self.assertEqual(_intents.output_request("search for pasta recipes and put it in a text file")["kind"], "txt")
        self.assertIsNone(_intents.output_request("open word")); self.assertIsNone(_intents.output_request("search youtube for lofi"))
    def test_safe_url(self):
        self.assertEqual(_intents.safe_url("https://www.forbes.com"), "https://www.forbes.com")
        self.assertEqual(_intents.safe_url("https://www.google.com/search?q=x"), "https://www.google.com/search?q=x")
        self.assertIn("google.com/search?q=site%3Aforbes.com+top+rich+list", _intents.safe_url("https://www.forbes.com/top-rich-list"))
        self.assertEqual(_intents.safe_url("https://www.forbes.com/real-time-billionaires/", ["https://www.forbes.com/real-time-billionaires/"]),
                         "https://www.forbes.com/real-time-billionaires/")
    def test_search_url_is_google(self):
        self.assertTrue(search_url("a b").startswith("https://www.google.com/search?q=a+b"))
    def test_make_documents(self):
        import docmaker, tempfile, zipfile
        rows = [[str(i), f"Person {i}", f"${300-i}B"] for i in range(1, 21)]
        with tempfile.TemporaryDirectory() as d:
            p = docmaker.make("docx", "Top 20 Richest", "Summary here.", ["Rank", "Name", "Net worth"], rows, "https://forbes.com/x", folder=d)
            from docx import Document
            doc = Document(p); t = doc.tables[0]
            self.assertEqual(len(t.rows), 21); self.assertEqual(t.rows[20].cells[1].text, "Person 20")
            body = "\n".join(x.text for x in doc.paragraphs); self.assertIn("Source: https://forbes.com/x", body); self.assertIn("Date:", body)
            x = docmaker.make("xlsx", "Top 20 Richest", "s", ["Rank", "Name"], rows, folder=d)
            import openpyxl; ws = openpyxl.load_workbook(x).active; self.assertEqual(ws.max_row, 4 + 1 + 20)
            t2 = docmaker.make("txt", "Top/20?", "s", ["Rank"], rows, folder=d); self.assertTrue(t2.endswith("Top20.txt"))
            self.assertTrue(p.endswith("Top 20 Richest.docx")); self.assertTrue(docmaker.make("docx", "Top 20 Richest", folder=d).endswith("(2).docx"))

class ForbesFake(FakeExt):
    def __init__(self, link, no_window=False):
        self.no_window = no_window; super().__init__(link)
    def snap(self):
        if "google" in self.url:
            els = [{"id": "1", "role": "link", "label": "Privacy", "href": "https://policies.google.com/privacy"},
                   {"id": "2", "role": "link", "label": "Forbes Real-Time Billionaires List", "href": "https://www.forbes.com/real-time-billionaires/"}]
        else: els = []
        return {"url": self.url, "title": "x", "elements": els, "dialogs": [], "text": "", "scroll": [0, 100], "status": 200}
    def loop(self):
        while not self.stop:
            for c in self.link.poll({"version": "test"}, hold=0.05):
                self.calls.append((c["cmd"], c["args"])); a = c["args"]; r = {"id": c["id"], "ok": True, "result": {}}
                if c["cmd"] == "snapshot": r["result"] = self.snap()
                elif c["cmd"] == "navigate":
                    if self.no_window: r = {"id": c["id"], "ok": False, "error": "no active tab"}; self.no_window = False
                    else: self.url = a["url"]
                elif c["cmd"] == "new_tab": self.url = a["url"]
                elif c["cmd"] == "extract_text":
                    r["result"] = {"text": " ".join(f"{i} Person{i} ${300-i}B" for i in range(1, 30)) if "forbes" in self.url else "results"}
                self.link.result(r)

class V127Agent(unittest.TestCase):
    def go(self, no_window=False):
        import extlink, time, tempfile, docmaker
        self.link = extlink.ExtensionLink(); self.ext = ForbesFake(self.link, no_window); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            time.sleep(0.01)
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 10, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        client = mock.MagicMock()
        rows = [[str(i), f"Person{i}", f"${300-i}B"] for i in range(1, 21)]
        client.text_chat.return_value = json.dumps({"title": "Top 20 Richest People", "summary": "Musk leads.", "columns": ["Rank", "Name", "Net worth"], "rows": rows})
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None
        d = tempfile.mkdtemp()
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"), mock.patch.object(docmaker, "out_dir", return_value=d), \
             mock.patch.object(docmaker, "open_file", return_value=True) as op:
            res = a.run(DOC_TASK)
        return res, client, op, d
    def test_user_sentence_makes_docx(self):
        res, client, op, d = self.go()
        self.assertIn("Saved 20 rows", res); self.assertIn("Person1, Person2", res)
        path = os.path.join(d, "Top 20 Richest People.docx"); self.assertTrue(os.path.exists(path)); op.assert_called_once_with(path)
        navs = [a["url"] for c, a in self.ext.calls if c in ("navigate", "new_tab")]
        self.assertEqual(navs, ["https://www.google.com/search?q=forbes+top+20+richest+people", "https://www.forbes.com/real-time-billionaires/"])
        client.chat.assert_not_called()  # never vision
    def test_no_window_opens_new_tab(self):
        res, client, op, d = self.go(no_window=True)
        self.assertIn("Saved 20 rows", res)
        self.assertIn("new_tab", [c for c, a in self.ext.calls])

XL_TASK = ("Open Excel, create a new sheet, and list the top 20 richest people in the world along with their net worth. "
           "Verify the information from a reliable source before finalizing the sheet.")
class V128AppFirst(unittest.TestCase):
    def test_exact_sentence(self):
        o = _intents.output_request(XL_TASK)
        self.assertEqual((o["kind"], o["research"]), ("xlsx", "list the top 20 richest people in the world along with their net worth"))
        self.assertEqual(_intents.research_query(o["research"]), (None, "top 20 richest people in the world along with their net worth"))
        self.assertEqual(_intents.requested_count(XL_TASK), 20); self.assertTrue(_intents.is_web_task(XL_TASK))
    def test_variants(self):
        cases = {"open word and write a summary of the latest iphone news": "docx",
                 "open notepad, make a new note and list the top 5 tallest buildings": "txt",
                 "can you open microsoft excel and make a list of the top 10 fastest cars": "xlsx",
                 "make an excel spreadsheet of the top 10 laptops of 2026": "xlsx",
                 "create a word document about the history of the internet": "docx",
                 "look up the top 10 richest people and put them in an excel sheet": "xlsx",
                 "find the latest apple news and save it in a word document": "docx",
                 "search for the top 5 tallest buildings and put it in a text file": "txt"}
        for t, k in cases.items():
            o = _intents.output_request(t); self.assertIsNotNone(o, t); self.assertEqual(o["kind"], k, t)
        for t in ("open excel", "open word and type hello", "open notepad", "open excel and create a new sheet"):
            self.assertIsNone(_intents.output_request(t), t)

class V128Agent(V127Agent):
    def test_excel_sentence_makes_xlsx(self):
        import extlink, time, tempfile, docmaker
        self.link = extlink.ExtensionLink(); self.ext = ForbesFake(self.link); self.addCleanup(setattr, self.ext, "stop", True)
        for _ in range(100):
            if self.link.connected(): break
            time.sleep(0.01)
        self.ext.snap_orig = self.ext.snap
        def snap():
            s = self.ext.snap_orig()
            if "google" in self.ext.url:
                s["elements"] = [{"id": "8", "role": "link", "label": "Some Blog Top Rich People", "href": "https://randomblog.example/rich"},
                                 {"id": "9", "role": "link", "label": "Forbes Real Time Billionaires", "href": "https://www.forbes.com/real-time-billionaires/"}]
            return s
        self.ext.snap = snap
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 10, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        client = mock.MagicMock()
        rows = [[f"Person{i}", f"${300-i}B"] for i in range(1, 21)]
        client.text_chat.return_value = json.dumps({"title": "Top 20 Richest", "summary": "s", "columns": ["Name", "Net Worth"], "rows": rows})
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: mock.MagicMock(), dom=False)
        a.kill.start = lambda: None; a.ext = self.link; a.focus_chrome = lambda: None
        d = tempfile.mkdtemp()
        with mock.patch.object(extlink, "LINK", self.link), mock.patch("time.sleep"), mock.patch.object(docmaker, "out_dir", return_value=d), \
             mock.patch.object(docmaker, "open_file", return_value=True) as op:
            res = a.run(XL_TASK)
        path = os.path.join(d, "Top 20 Richest.xlsx"); self.assertTrue(os.path.exists(path), res); op.assert_called_once_with(path)
        import openpyxl; ws = openpyxl.load_workbook(path).active
        vals = [r for r in ws.iter_rows(values_only=True)]
        self.assertIn("Source: https://www.forbes.com/real-time-billionaires/", [v[0] for v in vals])
        self.assertEqual(sum(1 for v in vals if v[0] and str(v[0]).startswith("Person")), 20)
        client.chat.assert_not_called()

class V128SameClick(unittest.TestCase):
    def test_same_click_three_times_stops(self):
        client = mock.MagicMock(); client.chat.return_value = '{"action":"click","x":676,"y":919}'
        g = mock.MagicMock(); g.size.return_value = (800, 600)
        cfg = {"ollama": {"host": "x", "model": "m"}, "agent": {"max_steps": 10, "runs_dir": "/tmp/pcagent-test", "step_delay_s": 0, "plan": False},
               "safety": {"confirm_risky": False, "kill_keys": []}}
        ex = mock.MagicMock()
        a = Agent(cfg, log=lambda m: None, gui=g, client=client, capture=lambda: Image.new("RGB", (800, 600)), executor=lambda s: ex, dom=False)
        a.kill.start = lambda: None
        with mock.patch.object(a, "use_extension", return_value=False), mock.patch("extlink.LINK.connected", return_value=False):
            res = a.run("click the cell in the spreadsheet app")
        self.assertIn("same spot (676,919) 3 times", res); self.assertEqual(client.chat.call_count, 3)
